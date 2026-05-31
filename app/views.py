from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import get_user_model, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth import authenticate, login as auth_login
from app.models import Request, RequestAttachment
from datetime import date, datetime
from django.db.models import Count
import openpyxl # For Excel export
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from django.http import HttpResponse
import io

User = get_user_model()

def login_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            if not user.is_approved and user.role != 'admin':
                return render(request, 'login.html', {
                    'error': 'Your account is pending admin approval.'
                })
            auth_login(request, user)
            return redirect('dashboard')
        else:
            return render(request, 'login.html', {'error': 'Invalid email or password.'})
    return render(request, 'login.html')

@login_required
def dashboard_view(request):
    user = request.user
    current_role = str(user.role).lower()

    if current_role == 'admin':
        tasks = Request.objects.all().order_by('-date_requested')

        total_count = tasks.count()
        completed_count = tasks.filter(status='Completed').count()
        ongoing_count = tasks.filter(status='Ongoing').count()
        requested_count = tasks.filter(status='Requested').count()

        recent_tasks = tasks[:5]

        categories = [
            'Graphic Design', 'Audio-Visual Presentation',
            'Photography and Editing', 'Web and Digital Design',
            'Publication Design', 'Branding and Identity',
        ]
        category_stats = []
        for cat in categories:
            cat_tasks = tasks.filter(category__iexact=cat)
            division_breakdown = (
                cat_tasks.values('division')
                .annotate(count=Count('id'))
                .order_by('-count')
            )
            category_stats.append({
                'name': cat,
                'count': cat_tasks.count(),
                'completed': cat_tasks.filter(status='Completed').count(),
                'ongoing': cat_tasks.filter(status='Ongoing').count(),
                'requested': cat_tasks.filter(status='Requested').count(),
                'in_progress': cat_tasks.filter(status='Ongoing').count(),
                'not_started': cat_tasks.filter(status='Requested').count(),
                'division_breakdown': division_breakdown,
                'requests': cat_tasks.order_by('-date_requested'),
            })

        # Pending users waiting for approval
        pending_users = User.objects.filter(is_approved=False, role__in=[
            'omcc', 'hopss', 'finance', 'allied', 'medical', 'nursing', 'staff'
        ]).order_by('-date_joined')

        # ── DIVISION STATS WITH MONTH FILTER ──
        selected_month = request.GET.get('month', 'all')

        if selected_month != 'all':
            try:
                month_date = datetime.strptime(selected_month, '%Y-%m')
                filtered_tasks = tasks.filter(
                    date_requested__year=month_date.year,
                    date_requested__month=month_date.month
                )
            except ValueError:
                filtered_tasks = tasks
        else:
            filtered_tasks = tasks

        # Get all distinct months that have requests (for the filter chips)
        available_months = (
            tasks.dates('date_requested', 'month', order='DESC')
        )

        # Build division stats using filtered tasks
        seen = set()
        divisions = tasks.values_list('division', flat=True).distinct()
        division_stats = []
        for div in divisions:
            if not div:
                continue
            div_clean = div.strip()  # remove any hidden spaces
            if div_clean in seen:    # skip duplicates
                continue
            seen.add(div_clean)
            div_tasks = filtered_tasks.filter(division__iexact=div_clean)
            division_stats.append({
                'name': div_clean,
                'count': div_tasks.count(),
                'completed': div_tasks.filter(status='Completed').count(),
                'ongoing': div_tasks.filter(status='Ongoing').count(),
                'requested': div_tasks.filter(status='Requested').count(),
                'requests': div_tasks.order_by('-date_requested'),
                'category_breakdown': (
                    div_tasks.values('category')
                    .annotate(count=Count('id'))
                    .order_by('-count')
                ),
            })

        context = {
            'tasks': tasks,
            'recent_tasks': recent_tasks,
            'category_stats': category_stats,
            'total_count': total_count,
            'completed_count': completed_count,
            'ongoing_count': ongoing_count,
            'requested_count': requested_count,
            'inprogress_count': tasks.filter(status='Ongoing').count(),
            'notstarted_count': tasks.filter(status='Requested').count(),
            'requested_tasks': tasks.filter(status='Requested').order_by('-date_requested'),
            'status_choices': ['Requested', 'Approved', 'Ongoing', 'Completed'],
            'staff_list': User.objects.filter(role='staff'),
            'pending_users': pending_users,
            'pending_count': pending_users.count(),
            'division_stats': division_stats,
            'available_months': available_months,
            'selected_month': selected_month,
        }
        return render(request, 'admin_dashboard.html', context)

    else:
        DIVISION_UNITS = {
            'hopss': [
                'Engineering & Facilities Management Section',
                'Biomed Unit', 'Building Facilities Unit',
                'Powerhouse Unit', 'Transportation Unit',
                'Human Resource Management Section',
                'Materials Management Section',
                'Procurement Section', 'General Services',
                'Housekeeping', 'Linen and Laundry',
            ],
            'allied': [
                'Health Information Management Section',
                'Admitting & Information Unit', 'OPD Records',
                'Nutrition & Dietetics Services',
                'Pharmacy Section', 'Social Service',
            ],
            'medical': [
                'Emergency Department', 'HEMS / DRRM-H',
                'Out Patient Services', 'Pathology & Laboratories',
                'Blood Bank', 'Molecular Laboratory', 'Paywards',
                'Radiology', 'Nuclear Medicine', 'Cancer Center',
                'Mental Health', 'Brain and Spine Center',
                'Cardiovascular Care', 'Eye Care',
                'Infectious Disease & Tropical Medicine',
                'Lung Care', 'Neonatal Care', 'Orthopedic Care',
                'Physical Rehabilitation Medicine',
                'Renal Care & Transplant Center', 'Trauma Care',
                'Anesthesiology', 'Family & Community Medicine',
                'Internal Medicine', 'Pediatrics',
                'Obstetrics & Gynecology',
                'Otorhinolaryngology - Head & Neck Surgery',
                'Surgery', 'Dental Medicine',
            ],
            'nursing': [
                'Amethyst Ward', 'Behavioral Medicine',
                'Medical Ward', 'OB-Gyne Ward', 'Pay Wards',
                'Pearl Ward', 'Pediatrics Ward', 'Surgery Ward',
                'Florence Nightingale', 'Diamond Ward', 'Garnet Ward',
                'Delivery Room', 'Intensive Care Unit', 'NICU',
                'PACU', 'PICU', 'SICU',
                'Frontline/Referring Services',
                'Specialty Centers', 'Clinics', 'HACT', 'TB DOTS',
            ],
            'finance': [
                'Accounting Section', 'Billing & Claims Section',
                'Budget Section', 'Cash Management Section',
            ],
            'omcc': ['DRRM-H', 'OSM', 'Quality Impovement Unit',
                     'Costumer Service', 'Legal Office',
                     'Planning & Management Unit',
                     'PETRU', 'PHU', 'IHOMP'],
        }

        user_units = DIVISION_UNITS.get(current_role, [])
        requests = Request.objects.filter(user=user).order_by('-date_requested')
        return render(request, 'user_dashboard.html', {
            'requests': requests,
            'all_requests_count': requests.count(),
            'pending_requests_count': requests.filter(status__in=['Requested', 'Ongoing', 'Approved']).count(),
            'completed_count': requests.filter(status='Completed').count(),
            'not_started_count': requests.filter(status='Requested').count(),
            'user_units': user_units,
        })


@login_required
def approve_user(request, pk):
    if request.user.role != 'admin':
        return redirect('dashboard')
    user = get_object_or_404(User, pk=pk)
    user.is_approved = True
    user.save()
    return redirect('dashboard')


@login_required
def reject_user(request, pk):
    if request.user.role != 'admin':
        return redirect('dashboard')
    user = get_object_or_404(User, pk=pk)
    user.delete()
    return redirect('dashboard')


@login_required
def update_task_status(request, pk):
    task = get_object_or_404(Request, pk=pk)
    if request.method == "POST" and str(request.user.role).lower() == 'admin':
        status = request.POST.get("status")
        staff_id = request.POST.get("assigned_to")
        if status:
            task.status = status
        if staff_id:
            try:
                task.assigned_to = User.objects.get(id=staff_id)
            except User.DoesNotExist:
                pass
        task.save()
    return redirect('dashboard')


@login_required
def create_request(request):
    if request.method == 'POST':
        target_date_raw = request.POST.get('target_date')
        target_date = date.fromisoformat(target_date_raw) if target_date_raw else None

        new_request = Request.objects.create(
            user=request.user,
            unit=request.POST.get('division'),
            project_title=request.POST.get('project_title'),
            category=request.POST.get('category'),
            request_type=request.POST.get('request_type'),
            description=request.POST.get('description'),
            target_date=target_date,
            delivery_type=request.POST.get('delivery_type') or None,
            delivery_detail=request.POST.get('delivery_detail') or None,
        )

        for f in request.FILES.getlist('attachments'):
            RequestAttachment.objects.create(request=new_request, file=f)

        return redirect('dashboard')
    return render(request, 'create_request.html')


@login_required
def update_request(request, pk):
    task = get_object_or_404(Request, pk=pk)
    if request.method == 'POST':
        task.project_title = request.POST.get('project_title', task.project_title)
        task.category = request.POST.get('category', task.category)
        task.request_type = request.POST.get('request_type', task.request_type)
        task.status = request.POST.get('status', task.status)
        task.division = request.POST.get('division', task.division)

        due_date_raw = request.POST.get('due_date')
        if due_date_raw:
            task.due_date = date.fromisoformat(due_date_raw)

        date_started_raw = request.POST.get('date_started')
        if date_started_raw:
            task.date_started = date.fromisoformat(date_started_raw)

        date_completed_raw = request.POST.get('date_completed')
        if date_completed_raw:
            task.date_completed = date.fromisoformat(date_completed_raw)

        staff_id = request.POST.get('assigned_to')
        if staff_id:
            try:
                task.assigned_to = User.objects.get(id=staff_id)
            except User.DoesNotExist:
                pass

        task.save()
        return redirect('/dashboard/?view=tasks')

    context = {
        'task': task,
        'staff_list': User.objects.filter(role='staff'),
        'status_choices': ['Requested', 'Approved', 'Ongoing', 'Completed'],
    }
    return render(request, 'update_request.html', context)


@login_required
def delete_request(request, pk):
    req = get_object_or_404(Request, pk=pk)
    if request.user.role == 'admin' or req.user == request.user:
        req.delete()
    return redirect('dashboard')


@login_required
def logout_view(request):
    logout(request)
    return redirect('login')


def register_view(request):
    if request.method == 'POST':
        full_name = request.POST.get('full_name')
        email = request.POST.get('email')
        role = request.POST.get('role')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        if password != confirm_password:
            return render(request, 'register.html', {'error': 'Passwords do not match'})

        if User.objects.filter(username=email).exists():
            return render(request, 'register.html', {'error': 'An account with this email already exists.'})
        
        if role.lower() == 'admin':
            admin_count = User.objects.filter(role='admin').count()
            if admin_count >= 2:
                return render(request, 'register.html', {'error': 'Maximum number of administrators has been reached.'})

        User.objects.create_user(
            username=email,
            email=email,
            password=password,
            full_name=full_name,
            role=role.lower(),
            is_approved=False,
        )
        return redirect('pending_approval')
    return render(request, 'register.html')

@login_required
def export_excel(request):
    if request.user.role != 'admin':
        return redirect('dashboard')

    export_type = request.GET.get('type', 'tasks')
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='0F1F3D', end_color='0F1F3D', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')

    tasks = Request.objects.all().order_by('-date_requested')

    if export_type == 'tasks' or export_type == 'all':
        ws = wb.active
        ws.title = 'All Tasks'
        headers = ['#', 'Date Requested', 'Division', 'Unit', 'Project Title',
                   'Type', 'Category', 'Target Date', 'Status', 'Staff Assigned',
                   'Delivery Type', 'Date Started', 'Date Completed']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for i, task in enumerate(tasks, 1):
            ws.append([
                i,
                task.date_requested.strftime('%b %d, %Y') if task.date_requested else '—',
                task.division or '—',
                task.unit or '—',
                task.project_title,
                task.request_type or '—',
                task.category or '—',
                task.target_date.strftime('%b %d, %Y') if task.target_date else '—',
                task.status,
                str(task.assigned_to) if task.assigned_to else '—',
                task.get_delivery_type_display() if task.delivery_type else '—',
                task.date_started.strftime('%b %d, %Y') if task.date_started else '—',
                task.date_completed.strftime('%b %d, %Y') if task.date_completed else '—',
            ])
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 4

    if export_type == 'category' or export_type == 'all':
        ws2 = wb.create_sheet('By Category')
        cat_headers = ['#', 'Category', 'Total', 'Completed', 'Ongoing', 'Requested']
        for col, h in enumerate(cat_headers, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        categories = ['Graphic Design', 'Audio-Visual Presentation',
                      'Photography and Editing', 'Web and Digital Design',
                      'Publication Design', 'Branding and Identity']
        for i, cat in enumerate(categories, 1):
            ct = tasks.filter(category__iexact=cat)
            ws2.append([i, cat, ct.count(),
                        ct.filter(status='Completed').count(),
                        ct.filter(status='Ongoing').count(),
                        ct.filter(status='Requested').count()])
        for col in ws2.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws2.column_dimensions[col[0].column_letter].width = max_len + 4

    if export_type == 'division' or export_type == 'all':
        ws3 = wb.create_sheet('By Division')
        div_headers = ['#', 'Division', 'Total', 'Completed', 'Ongoing', 'Requested']
        for col, h in enumerate(div_headers, 1):
            cell = ws3.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        divisions = tasks.values_list('division', flat=True).distinct()
        for i, div in enumerate(divisions, 1):
            if not div:
                continue
            dt = tasks.filter(division=div)
            ws3.append([i, div, dt.count(),
                        dt.filter(status='Completed').count(),
                        dt.filter(status='Ongoing').count(),
                        dt.filter(status='Requested').count()])
        for col in ws3.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws3.column_dimensions[col[0].column_letter].width = max_len + 4

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'CVMC_Multimedia_Report_{export_type}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required # Only admins can export PDF reports
def export_pdf(request):
    if request.user.role != 'admin':
        return redirect('dashboard')

    export_type = request.GET.get('type', 'tasks')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            rightMargin=20, leftMargin=20,
                            topMargin=30, bottomMargin=20)
    styles = getSampleStyleSheet()
    elements = []

    header_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F1F3D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f4fa')]),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWHEIGHT', (0, 0), (-1, -1), 20),
    ])

    tasks = Request.objects.all().order_by('-date_requested')

    if export_type == 'tasks' or export_type == 'all':
        elements.append(Paragraph('All Tasks Report', styles['Title']))
        elements.append(Paragraph('CVMC Multimedia Task Management Tracker', styles['Normal']))
        elements.append(Spacer(1, 12))
        data = [['#', 'Date', 'Division', 'Unit', 'Project', 'Type', 'Category', 'Target', 'Status', 'Staff']]
        for i, task in enumerate(tasks, 1):
            data.append([
                str(i),
                task.date_requested.strftime('%b %d') if task.date_requested else '—',
                task.division or '—',
                (task.unit or '—')[:15],
                task.project_title[:20],
                (task.request_type or '—')[:12],
                (task.category or '—')[:15],
                task.target_date.strftime('%b %d, %Y') if task.target_date else '—',
                task.status,
                str(task.assigned_to)[:15] if task.assigned_to else '—',
            ])
        t = Table(data, repeatRows=1)
        t.setStyle(header_style)
        elements.append(t)

    if export_type == 'category' or export_type == 'all':
        if export_type == 'all':
            elements.append(Spacer(1, 24))
        elements.append(Paragraph('By Category Report', styles['Title']))
        elements.append(Spacer(1, 12))
        categories = ['Graphic Design', 'Audio-Visual Presentation',
                      'Photography and Editing', 'Web and Digital Design',
                      'Publication Design', 'Branding and Identity']
        data2 = [['#', 'Category', 'Total', 'Completed', 'Ongoing', 'Requested']]
        for i, cat in enumerate(categories, 1):
            ct = tasks.filter(category__iexact=cat)
            data2.append([str(i), cat, str(ct.count()),
                          str(ct.filter(status='Completed').count()),
                          str(ct.filter(status='Ongoing').count()),
                          str(ct.filter(status='Requested').count())])
        t2 = Table(data2, repeatRows=1)
        t2.setStyle(header_style)
        elements.append(t2)

    if export_type == 'division' or export_type == 'all':
        if export_type == 'all':
            elements.append(Spacer(1, 24))
        elements.append(Paragraph('By Division Report', styles['Title']))
        elements.append(Spacer(1, 12))
        divisions = tasks.values_list('division', flat=True).distinct()
        data3 = [['#', 'Division', 'Total', 'Completed', 'Ongoing', 'Requested']]
        for i, div in enumerate(divisions, 1):
            if not div:
                continue
            dt = tasks.filter(division=div)
            data3.append([str(i), div, str(dt.count()),
                          str(dt.filter(status='Completed').count()),
                          str(dt.filter(status='Ongoing').count()),
                          str(dt.filter(status='Requested').count())])
        t3 = Table(data3, repeatRows=1)
        t3.setStyle(header_style)
        elements.append(t3)

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    filename = f'CVMC_Multimedia_Report_{export_type}.pdf'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

def forgot_password_view(request):
    return render(request, 'forgot_password.html')

def pending_approval_view(request):
    return render(request, 'pending_approval.html')